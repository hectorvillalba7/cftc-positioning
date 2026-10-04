#!/usr/bin/env python3
"""Compara los datos calculados (docs/data) con tu Excel manual.

Sirve para comprobar, tras la primera carga, que la web da exactamente los
mismos números que apuntabas a mano (códigos de contrato, inversión de
largos y cortos, solo futuros o futuros + opciones).

    pip install openpyxl
    python tools/compare_excel.py CFTC_DATA_Tiff.xlsx            # reporte tff
    python tools/compare_excel.py Legacy.xlsx --report legacy

Lee la cabecera de la fila 3 de cada hoja (un bloque por grupo, que empieza
en cada columna "Date"), así que da igual que una hoja tenga o no las
columnas de Total o de % de traders. Los datos se leen desde la fila 7.
"""

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent


def parse_layout(ws, header_row=3):
    """Devuelve, por cada bloque, {campo: índice de columna} según la cabecera."""
    labels = [str(c.value).replace(" ", "").lower() if c.value is not None else "" for c in ws[header_row]]
    starts = [i for i, v in enumerate(labels) if v == "date"]
    blocks = []
    for n, start in enumerate(starts):
        end = starts[n + 1] if n + 1 < len(starts) else len(labels)
        longs = [i for i in range(start + 1, end) if labels[i] == "long"]
        pcts = [i for i in range(start + 1, end) if labels[i] == "%long"]
        block = {"date": start}
        if len(longs) == 2:      # traders y posiciones
            block.update(tl=longs[0], ts=longs[0] + 1, pl=longs[1], ps=longs[1] + 1)
        elif len(longs) == 1:    # solo posiciones
            block.update(pl=longs[0], ps=longs[0] + 1)
        if pcts:                 # el último par de % es el de open interest
            block.update(pol=pcts[-1], pos=pcts[-1] + 1)
        blocks.append(block)
    return blocks

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--report", default="tff")
    ap.add_argument("--data", default=str(ROOT / "docs" / "data"))
    args = ap.parse_args()

    data_dir = Path(args.data)
    index = json.loads((data_dir / "index.json").read_text(encoding="utf-8"))
    report = next(r for r in index["reports"] if r["key"] == args.report)
    wb = load_workbook(args.excel, data_only=True)
    sheets = {name.replace(" ", "").lower(): name for name in wb.sheetnames}

    total_bad = 0
    for asset in report["assets"]:
        sheet = sheets.get(asset["name"].replace(" ", "").lower())
        if not sheet:
            print(f"{asset['name']}: no hay hoja con ese nombre en el Excel")
            continue
        data = json.loads((data_dir / args.report / f"{asset['id']}.json").read_text(encoding="utf-8"))
        by_date = dict(zip(data["dates"], data["rows"]))
        col = {c: i for i, c in enumerate(data["columns"])}
        layout = parse_layout(wb[sheet])
        checked, bad = 0, []
        for row in wb[sheet].iter_rows(min_row=7, values_only=True):
            if row[0] is None or row[1] is None:
                continue
            d = row[1].strftime("%Y-%m-%d")
            web = by_date.get(d)
            if web is None:
                bad.append(f"  {d}: la fecha no existe en los datos de la CFTC")
                continue
            pairs = [("total traders", row[0], web[col["tt"]])]
            for g, block in zip(report["groups"], layout):
                for f in ("tl", "ts", "pl", "ps", "pol", "pos"):
                    if f in block and block[f] < len(row):
                        pairs.append((f"{g['label']} {f}", row[block[f]], web[col[f"{g['key']}.{f}"]]))
            for label, excel_v, web_v in pairs:
                if excel_v is None:
                    continue
                checked += 1
                if web_v is None or abs(float(excel_v) - float(web_v)) > 0.051:
                    bad.append(f"  {d} {label}: Excel {excel_v} / web {web_v}")
        total_bad += len(bad)
        print(f"{asset['name']}: {checked} celdas comparadas, {len(bad)} diferencias")
        for line in bad[:8]:
            print(line)
        if len(bad) > 8:
            print(f"  ... y {len(bad) - 8} más")
    sys.exit(1 if total_bad else 0)


if __name__ == "__main__":
    main()
