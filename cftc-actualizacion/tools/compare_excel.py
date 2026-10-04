#!/usr/bin/env python3
"""Compara los datos calculados (docs/data) con tus Excel manuales.

Sirve para comprobar, tras una carga, que la web da exactamente los mismos
números que apuntabas a mano (códigos de contrato, inversión de largos y
cortos, solo futuros o futuros + opciones).

    pip install openpyxl
    python tools/compare_excel.py CFTC_DATA_Tiff.xlsx
    python tools/compare_excel.py CFTC_DATA_Legacy.xlsx --report legacy
    python tools/compare_excel.py CFTC_DATA_Disagreggated.xlsx --report disaggregated

Lee la cabecera de cada hoja (la fila donde están las columnas "Date"): cada
"Date" abre el bloque de un grupo, en el mismo orden que `groups` en
config.yaml. Da igual que una hoja tenga o no las columnas de total de
traders, Total o % de traders.
"""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent


def norm(value):
    return str(value).replace(" ", "").replace("_", "").lower() if value is not None else ""


def parse_layout(ws):
    """Devuelve (fila de cabecera, columna del total de traders o None, bloques).

    Cada bloque es {campo: índice de columna}.
    """
    for header_row in range(1, 9):
        labels = [norm(c.value) for c in ws[header_row]]
        if "date" in labels:
            break
    else:
        return None
    total_col = next((i for i, v in enumerate(labels) if v in ("ttraders", "totaltraders")), None)
    starts = [i for i, v in enumerate(labels) if v == "date"]
    blocks = []
    for n, start in enumerate(starts):
        end = starts[n + 1] if n + 1 < len(starts) else len(labels)
        longs = [i for i in range(start + 1, end) if labels[i] == "long"]
        pcts = [i for i in range(start + 1, end) if labels[i] == "%long"]
        block = {"date": start}
        if len(longs) == 2:      # traders y posiciones
            block.update(tl=longs[0], ts=longs[0] + 1, pl=longs[1], ps=longs[1] + 1)
        elif len(longs) == 1:    # solo posiciones
            block.update(pl=longs[0], ps=longs[0] + 1)
        if pcts:                 # el último par de % es el de open interest
            block.update(pol=pcts[-1], pos=pcts[-1] + 1)
        blocks.append(block)
    return header_row, total_col, blocks


def find_sheet(wb, asset):
    sheets = {norm(name): name for name in wb.sheetnames}
    for key in (norm(asset["name"]), norm(asset["id"]), norm(asset["name"]).rstrip("s")):
        if key in sheets:
            return sheets[key]
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--report", default="tff")
    ap.add_argument("--data", default=str(ROOT / "docs" / "data"))
    args = ap.parse_args()

    data_dir = Path(args.data)
    index = json.loads((data_dir / "index.json").read_text(encoding="utf-8"))
    report = next(r for r in index["reports"] if r["key"] == args.report)
    wb = load_workbook(args.excel, data_only=True)

    total_bad = 0
    for asset in report["assets"]:
        sheet = find_sheet(wb, asset)
        layout = parse_layout(wb[sheet]) if sheet else None
        if not layout:
            print(f"{asset['name']}: no hay hoja con ese nombre en el Excel")
            continue
        header_row, total_col, blocks = layout
        data = json.loads((data_dir / args.report / f"{asset['id']}.json").read_text(encoding="utf-8"))
        by_date = dict(zip(data["dates"], data["rows"]))
        col = {c: i for i, c in enumerate(data["columns"])}
        checked, bad = 0, []
        for row in wb[sheet].iter_rows(min_row=header_row + 1, values_only=True):
            date = row[blocks[0]["date"]]
            if not isinstance(date, (dt.datetime, dt.date)):
                continue
            d = date.strftime("%Y-%m-%d")
            pairs = []
            if total_col is not None:
                pairs.append(("total traders", row[total_col], "tt"))
            for g, block in zip(report["groups"], blocks):
                for f in ("tl", "ts", "pl", "ps", "pol", "pos"):
                    if f in block and block[f] < len(row):
                        pairs.append((f"{g['label']} {f}", row[block[f]], f"{g['key']}.{f}"))
            pairs = [(label, v, c) for label, v, c in pairs if isinstance(v, (int, float))]
            if not pairs:
                continue
            web = by_date.get(d)
            if web is None:
                bad.append(f"  {d}: la fecha no existe en los datos de la CFTC")
                continue
            for label, excel_v, c in pairs:
                checked += 1
                web_v = web[col[c]]
                if web_v is None or abs(float(excel_v) - float(web_v)) > 0.051:
                    bad.append(f"  {d} {label}: Excel {excel_v} / web {web_v}")
        total_bad += len(bad)
        print(f"{asset['name']}: {checked} celdas comparadas, {len(bad)} diferencias")
        for line in bad[:8]:
            print(line)
        if len(bad) > 8:
            print(f"  ... y {len(bad) - 8} más")
    sys.exit(1 if total_bad else 0)


if __name__ == "__main__":
    main()
